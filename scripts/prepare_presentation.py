"""Fill and enhance the supplied Samsung presentation template."""
import argparse
from pathlib import Path
import zipfile
import xml.etree.ElementTree as ET

NS = {'a': 'http://schemas.openxmlformats.org/drawingml/2006/main',
      'p': 'http://schemas.openxmlformats.org/presentationml/2006/main'}
for prefix, uri in NS.items():
    ET.register_namespace(prefix, uri)

CONTENT = {
    1: (5, [
        'Theme 01 — Agentic Code Intelligence',
        'Project — CodeStrata: code has a memory',
        'Team — CodeStrata',
        'Vellore Institute of Technology (VIT University)',
        'Advay Vivek · 24BCT0136 · advay.vivek2024@vitstudent.ac.in',
        'Kaavish Gogia · 24BCT0103 · kaavish.gogia2024@vitstudent.ac.in',
        'GitHub — github.com/AdvayV/AgenticIntelligence',
    ]),
    2: (1, [
        'Find the exact version where behavior changed.',
        'Even when the function was renamed and moved.',
        'CPU-only JavaScript retrieval with source, lines, versions, and evidence.',
        'P0: retrieval quality. P1: versions. Bonus: evolutionary investigation.',
        'Output is reviewable code evidence; no generated fixes.',
    ]),
    3: (1, [
        'Similar snippets can behave differently after a one-token change.',
        'Names and file paths stop being reliable after refactors.',
        'A ranked match alone does not explain a wrong-version result.',
        'Our focus: contrast the match with a version that contradicts it.',
        'Existing code search and agentic investigation already exist.',
        'Differentiation is the integration; no universal novelty claim.',
    ]),
    4: (1, []),
    5: (1, []),
    6: (1, [
        'Node.js + JavaScript; Babel parser; plain HTML/CSS/modules.',
        'BM25 + normalized cosine vectors + reciprocal rank fusion.',
        'Optional Transformers.js q8 BGE / MiniLM on CPU, two threads.',
        'Deterministic planner; no LLM API, GPU, or repository execution.',
        'Python/MTEB for official evaluation; persistent Node encoder.',
        'Node tests, jsdom, Playwright, Docker, GitHub Actions.',
    ]),
    7: (1, [
        'Review a refactor that silently removed a direct await.',
        'Locate changed guard presence or simple invocation order.',
        'Find an exact deeplink and inspect its version-specific source.',
        'Navigate related helpers without inventing execution guarantees.',
        'Useful for regression triage and version migration reviews.',
        'Static evidence supports inspection; it does not prove a runtime bug.',
    ]),
    8: (1, []),
    9: (1, [
        'Prioritize a code-specific encoder using separate development data.',
        'Collect independently labeled multi-commit repository questions.',
        'Explain the gap between development and official test accuracy.',
        'Broaden conservative lineage and supported import patterns.',
        'Measure larger repositories before adding approximate indexes.',
        'Finish team identity, video, disclosure, and organizer review.',
    ]),
    10: (1, [
        'Behavior evidence survives a simultaneous rename and file move.',
        'Counterexamples expose why a nearby version does not match.',
        'Ambiguous copies and refactors stay unresolved.',
        'Live agent events, not a simulated progress animation.',
        'Light timeline workspace with hover/focus/tap explanations.',
        'CPU-only, transparent limits, and reproducible raw results.',
    ]),
    11: (1, [
        'Prototype code + setup README — YES, public GitHub.',
        'Tests + browser/Docker workflows — see latest Actions status.',
        'Official MTEB JSON + real-repository artifacts — YES, docs/results.',
        'Presentation draft — YES; team details remain incomplete.',
        'Video + signed AI disclosure — PENDING team action.',
        'Final judged release/tag — PENDING submission readiness.',
    ]),
}


def fill(xml, shape_index, lines, size, number):
    root = ET.fromstring(xml)
    shape = root.findall('.//p:sp', NS)[shape_index]
    body = shape.find('p:txBody', NS)
    if body is None:
        body = ET.SubElement(shape, '{' + NS['p'] + '}txBody')
        ET.SubElement(body, '{' + NS['a'] + '}bodyPr')
        ET.SubElement(body, '{' + NS['a'] + '}lstStyle')
    for paragraph in list(body.findall('a:p', NS)):
        body.remove(paragraph)
    for line in lines:
        paragraph = ET.SubElement(body, '{' + NS['a'] + '}p')
        props = ET.SubElement(paragraph, '{' + NS['a'] + '}pPr')
        ET.SubElement(props, '{' + NS['a'] + '}buNone')
        spacing = ET.SubElement(props, '{' + NS['a'] + '}spcAft')
        ET.SubElement(spacing, '{' + NS['a'] + '}spcPts', {'val': '1200'})
        run = ET.SubElement(paragraph, '{' + NS['a'] + '}r')
        style = ET.SubElement(run, '{' + NS['a'] + '}rPr', {'lang': 'en-US', 'sz': str(size)})
        color = ET.SubElement(style, '{' + NS['a'] + '}solidFill')
        ET.SubElement(color, '{' + NS['a'] + '}srgbClr', {'val': '24243B'})
        ET.SubElement(style, '{' + NS['a'] + '}latin', {'typeface': 'Aptos'})
        ET.SubElement(run, '{' + NS['a'] + '}t').text = line
    if shape_index == 5:
        # The supplied title-slide placeholder is only 0.85 in high.
        # Expand it to keep both contributors visible below the title.
        shape.find('.//a:xfrm/a:ext', NS).set('cy', '2743200')
    if number == 4:
        add_architecture(root)
    return ET.tostring(root, encoding='utf-8', xml_declaration=True)


def add_architecture(root):
    tree = root.find('.//p:spTree', NS)
    existing = [int(item.get('id')) for item in root.findall('.//p:cNvPr', NS)]
    next_id = max(existing) + 1

    def box(label, x, y, width, height, fill, size=1800):
        nonlocal next_id
        shape = ET.SubElement(tree, '{' + NS['p'] + '}sp')
        nv = ET.SubElement(shape, '{' + NS['p'] + '}nvSpPr')
        ET.SubElement(nv, '{' + NS['p'] + '}cNvPr', {'id': str(next_id), 'name': f'CodeStrata diagram {next_id}'})
        next_id += 1
        ET.SubElement(nv, '{' + NS['p'] + '}cNvSpPr')
        ET.SubElement(nv, '{' + NS['p'] + '}nvPr')
        sppr = ET.SubElement(shape, '{' + NS['p'] + '}spPr')
        transform = ET.SubElement(sppr, '{' + NS['a'] + '}xfrm')
        ET.SubElement(transform, '{' + NS['a'] + '}off', {'x': str(x), 'y': str(y)})
        ET.SubElement(transform, '{' + NS['a'] + '}ext', {'cx': str(width), 'cy': str(height)})
        geom = ET.SubElement(sppr, '{' + NS['a'] + '}prstGeom', {'prst': 'roundRect'})
        ET.SubElement(geom, '{' + NS['a'] + '}avLst')
        solid = ET.SubElement(sppr, '{' + NS['a'] + '}solidFill')
        ET.SubElement(solid, '{' + NS['a'] + '}srgbClr', {'val': fill})
        line = ET.SubElement(sppr, '{' + NS['a'] + '}ln', {'w': '12700'})
        linefill = ET.SubElement(line, '{' + NS['a'] + '}solidFill')
        ET.SubElement(linefill, '{' + NS['a'] + '}srgbClr', {'val': 'D9D3EE'})
        body = ET.SubElement(shape, '{' + NS['p'] + '}txBody')
        ET.SubElement(body, '{' + NS['a'] + '}bodyPr', {'anchor': 'ctr', 'lIns': '110000', 'rIns': '110000'})
        ET.SubElement(body, '{' + NS['a'] + '}lstStyle')
        paragraph = ET.SubElement(body, '{' + NS['a'] + '}p')
        ppr = ET.SubElement(paragraph, '{' + NS['a'] + '}pPr', {'algn': 'ctr'})
        ET.SubElement(ppr, '{' + NS['a'] + '}buNone')
        run = ET.SubElement(paragraph, '{' + NS['a'] + '}r')
        rpr = ET.SubElement(run, '{' + NS['a'] + '}rPr', {'lang': 'en-US', 'sz': str(size), 'b': '1'})
        textfill = ET.SubElement(rpr, '{' + NS['a'] + '}solidFill')
        ET.SubElement(textfill, '{' + NS['a'] + '}srgbClr', {'val': '33255B'})
        ET.SubElement(run, '{' + NS['a'] + '}t').text = label

    # The diagram stays within the template's content rectangle.
    x_positions = (820000, 4440000, 8060000)
    for x, label in zip(x_positions, ('Git versions', 'AST behavior facts', 'BM25 + CPU vectors')):
        box(label, x, 2000000, 2850000, 820000, 'F1EEFF')
    for x, label in zip(x_positions, ('Question + constraints', 'Bounded search agent', 'Source + counterexample')):
        box(label, x, 3890000, 2850000, 820000, 'EAF8F2')
    for x in (3670000, 7290000):
        box('→', x, 2130000, 570000, 530000, 'FFFFFF', 2200)
        box('→', x, 4020000, 570000, 530000, 'FFFFFF', 2200)
    box('Lineage + import graph  ↓  evidence inspection', 820000, 3060000, 10090000, 530000, 'FFFFFF', 1400)
    box('Conservative static evidence · live agent events · CPU only', 820000, 5160000, 10090000, 480000, 'F9F7FD', 1400)


def enhance_deck(target):
    """Use vector cards so the demo and benchmark slides are readable in a room."""
    from pptx import Presentation
    from pptx.dml.color import RGBColor
    from pptx.enum.shapes import MSO_SHAPE
    from pptx.enum.text import PP_ALIGN, MSO_ANCHOR
    from pptx.util import Inches, Pt

    deck = Presentation(target)
    ink, purple, mint, amber = (RGBColor.from_string(s) for s in ('24243B', '6546CF', 'EAF8F2', 'FFF2E5'))

    def panel(slide, x, y, width, height, fill, border='E2DDEF'):
        shape = slide.shapes.add_shape(MSO_SHAPE.ROUNDED_RECTANGLE, Inches(x), Inches(y), Inches(width), Inches(height))
        shape.fill.solid()
        shape.fill.fore_color.rgb = RGBColor.from_string(fill)
        shape.line.color.rgb = RGBColor.from_string(border)
        return shape

    def label(slide, text, x, y, width, height, size=15, color=ink, bold=False, font='Aptos', align=None):
        box = slide.shapes.add_textbox(Inches(x), Inches(y), Inches(width), Inches(height))
        frame = box.text_frame
        frame.clear()
        frame.word_wrap = True
        frame.margin_left = frame.margin_right = Inches(0.04)
        frame.margin_top = frame.margin_bottom = Inches(0.02)
        frame.vertical_anchor = MSO_ANCHOR.MIDDLE
        paragraph = frame.paragraphs[0]
        if align:
            paragraph.alignment = align
        run = paragraph.add_run()
        run.text = text
        run.font.name = font
        run.font.size = Pt(size)
        run.font.bold = bold
        run.font.color.rgb = color
        return box

    slide = deck.slides[4]
    label(slide, 'WHERE DID DEVICE PAIRING STOP WAITING FOR policyCheck?', .91, 1.83, 11.2, .37, 17, purple, True)
    cards = [
        ('v1 · connectDevice', 'await policyCheck()', 'devices/connect.js', 'F0EBFF'),
        ('v2 · pairDevice', 'policyCheck()', 'flows/pair.js', 'FFF2E5'),
        ('v3 · pairDevice', 'await policyCheck()', 'flows/pair.js', 'EAF8F2'),
    ]
    for idx, (title, code, location, fill) in enumerate(cards):
        x = .91 + idx * 4.18
        panel(slide, x, 2.33, 3.7, 1.08, fill)
        label(slide, title, x + .18, 2.47, 3.32, .28, 17, ink, True)
        label(slide, code + '  ·  ' + location, x + .18, 2.88, 3.35, .32, 12, ink, False, 'Consolas')
        if idx < 2:
            label(slide, '→', x + 3.79, 2.64, .32, .38, 20, purple, True)
    panel(slide, .91, 3.76, 5.52, 1.47, 'F8F6FD')
    panel(slide, 6.77, 3.76, 5.52, 1.47, 'F8F6FD')
    label(slide, 'MATCH · v2', 1.11, 3.89, 4.9, .28, 13, purple, True)
    label(slide, 'policyCheck(device);\nreturn openDevice(device);', 1.11, 4.27, 5.0, .65, 17, ink, False, 'Consolas')
    label(slide, 'COUNTEREXAMPLE · v1', 6.97, 3.89, 4.9, .28, 13, RGBColor.from_string('96562E'), True)
    label(slide, 'await policyCheck(device);\nreturn openDevice(device);', 6.97, 4.27, 5.0, .65, 17, ink, False, 'Consolas')
    label(slide, 'Click the timeline to inspect exact source. Follow the imported policyCheck helper. Static evidence does not prove a runtime race.', .91, 5.5, 11.5, .7, 17, ink)
    label(slide, 'Video URL: [TEAM TO ADD AFTER RECORDING · 5 MIN MAX]', .91, 6.24, 11.5, .36, 13, purple, True)

    slide = deck.slides[7]
    metrics = [
        ('CONTROLLED VERSIONS', '8 / 8', 'correct top result', '25 demo snippets + 1,000 distractors', 'F0EBFF'),
        ('REAL CODE', '0.67 / 0.86', 'NDCG@10, Async / Express', '60 AI-assisted, source-reviewed labels', 'EAF8F2'),
        ('OFFICIAL SCREENING', '0.0505', 'AppsRetrieval NDCG@10', 'weak baseline · improvement required', 'FFF2E5'),
    ]
    for idx, (title, value, detail, footnote, fill) in enumerate(metrics):
        x = .9 + idx * 4.18
        panel(slide, x, 2.0, 3.79, 2.75, fill)
        label(slide, title, x + .18, 2.16, 3.42, .32, 13, purple, True)
        label(slide, value, x + .18, 2.55, 3.42, .76, 36, ink, True)
        label(slide, detail, x + .18, 3.43, 3.42, .43, 16)
        label(slide, footnote, x + .18, 4.08, 3.42, .48, 11, ink)
    label(slide, 'These are separate evaluations. Official MRR@10 is 0.043363. The MTEB encoder does not test version reranking.', .91, 5.12, 11.5, .66, 17, ink)
    label(slide, 'Limits: heuristic lineage · bounded static syntax · no independently annotated real-code labels.', .91, 5.97, 11.5, .5, 14, purple, True)
    deck.save(target)


def main():
    parser = argparse.ArgumentParser()
    parser.add_argument('--template', required=True)
    parser.add_argument('--output', default='docs/CodeStrata-Submission-Draft.pptx')
    args = parser.parse_args()
    target = Path(args.output)
    target.parent.mkdir(parents=True, exist_ok=True)
    with zipfile.ZipFile(args.template) as source, zipfile.ZipFile(target, 'w', zipfile.ZIP_DEFLATED) as out:
        for entry in source.infolist():
            payload = source.read(entry.filename)
            for number, (shape, lines) in CONTENT.items():
                if entry.filename == f'ppt/slides/slide{number}.xml':
                    payload = fill(payload, shape, lines, 1600 if number == 1 else 2200, number)
            out.writestr(entry, payload)
    with zipfile.ZipFile(target) as check:
        assert check.testzip() is None
        for name in check.namelist():
            if name.endswith('.xml'):
                ET.fromstring(check.read(name))
    enhance_deck(target)
    with zipfile.ZipFile(target) as check:
        assert check.testzip() is None
        for name in check.namelist():
            if name.endswith('.xml'):
                ET.fromstring(check.read(name))
    print(f'Prepared and XML-validated 12-slide draft: {target}')


if __name__ == '__main__':
    main()
